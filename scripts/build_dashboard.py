#!/usr/bin/env python3
"""Build the Pool Pathology downloads dashboard (index.html) from the cumulative Excel report.

Usage:
    python scripts/build_dashboard.py path/to/REPORTCumulative_Pool_Pathology_Analytics_<dates>.xlsx [-o index.html]

What it reads from the workbook:
    AllDownloads              every download: Package ID, Package Name, Date
    PkgID-Name-Count          package list (cols A-C) and the category totals (cols E-F);
                              each category's SUM(...) formula defines which packages belong to it
    Download Number by Date   column C, "E-mail Sent Dates"

What it writes:
    One self-contained HTML file: page, styles, D3 and all data inlined. No server needed.
"""
import argparse
import datetime as dt
import html
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "template" / "dashboard_template.html"
APP_JS = ROOT / "template" / "app.js"
D3 = ROOT / "vendor" / "d3.v7.9.0.min.js"


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):  # Excel serial
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    return dt.date.fromisoformat(str(v)[:10])


def expand(ref_list):
    """'C5,C20:C22' -> [5, 20, 21, 22]"""
    rows = []
    for part in ref_list.split(","):
        a, _, b = part.strip().partition(":")
        a = int(re.sub(r"\D", "", a))
        b = int(re.sub(r"\D", "", b)) if b else a
        rows.extend(range(a, b + 1))
    return rows


def load(path):
    wb = openpyxl.load_workbook(path)  # formulas kept, so category ranges are readable

    # packages + categories
    ws = wb["PkgID-Name-Count"]
    row_cat = {}
    for r in range(2, ws.max_row + 1):
        label, f = ws.cell(r, 5).value, ws.cell(r, 6).value
        if not label or str(label).strip().upper() == "SUM":
            break
        m = re.match(r"=SUM\((.*)\)", str(f or ""))
        if not m:
            continue
        for row in expand(m.group(1)):
            row_cat[row] = str(label).strip()

    packages = []
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, 1).value
        if pid is None:
            continue
        name = ws.cell(r, 2).value
        name = html.unescape(str(name)) if name else f"(Unnamed package {int(pid)})"
        packages.append({
            "package_id": int(pid),
            "package": name,
            "category": row_cat.get(r, "Uncategorised"),
        })

    # downloads
    counts = Counter()
    for pid, _name, day, *_ in wb["AllDownloads"].iter_rows(min_row=2, values_only=True):
        if pid is None or day is None:
            continue
        counts[(as_date(day).isoformat(), int(pid))] += 1

    known = {p["package_id"] for p in packages}
    for _, pid in counts:
        if pid not in known:  # a package downloaded but missing from the package tab
            known.add(pid)
            packages.append({"package_id": pid, "package": f"(Package {pid})", "category": "Uncategorised"})

    # e-mail send dates
    emails = sorted({
        as_date(v).isoformat()
        for (v,) in wb["Download Number by Date"].iter_rows(min_row=2, min_col=3, max_col=3, values_only=True)
        if v is not None and not isinstance(v, str)
    })
    return packages, counts, emails


def build(src, out):
    packages, counts, emails = load(src)
    days = sorted({d for d, _ in counts})
    idx = {d: i for i, d in enumerate(days)}
    data = {
        "days": days,
        "rows": [[idx[d], pid, n] for (d, pid), n in sorted(counts.items())],
        "packages": packages,
        "emails": [{"day": d, "event": "Bone Path (BP) e-mail sent"} for d in emails],
    }
    fmt = lambda s: dt.date.fromisoformat(s).strftime("%-m/%-d/%Y")
    page = TEMPLATE.read_text(encoding="utf-8")
    page = page.replace("/*__D3__*/", D3.read_text(encoding="utf-8"))
    page = page.replace("/*__DATA__*/", json.dumps(data, separators=(",", ":"), ensure_ascii=False))
    page = page.replace("/*__APP__*/", APP_JS.read_text(encoding="utf-8"))
    page = page.replace("__FIRST_DAY__", fmt(days[0])).replace("__LAST_DAY__", fmt(days[-1]))
    Path(out).write_text(page, encoding="utf-8")
    total = sum(counts.values())
    print(f"Wrote {out}: {total:,} downloads, {len(packages)} packages, {len(days)} active days "
          f"({days[0]} to {days[-1]}), {len(emails)} e-mail dates")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("report", help="Cumulative Pool Pathology report (.xlsx)")
    ap.add_argument("-o", "--out", default=str(ROOT / "index.html"), help="Output HTML (default: index.html)")
    a = ap.parse_args()
    build(a.report, a.out)
